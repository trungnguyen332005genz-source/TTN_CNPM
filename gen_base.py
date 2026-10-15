# -*- coding: utf-8 -*-
"""
Full Generator for KeHoachCNPM.docx
Strictly focusing on Nguyen's assigned part: Mô hình hóa luồng xử lý
Luân's part is kept intact as requested.
"""
import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)
    r = h.runs[0]
    r.font.name = 'Times New Roman'
    if level == 1:
        r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(31, 78, 121) # deep blue
    elif level == 2:
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(46, 117, 182)
    elif level == 3:
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(68, 114, 196)
    return h

def add_p(doc, text="", bold=False, italic=False, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    c = table.cell(0, 0)
    c.width = Inches(6.5)
    set_cell_background(c, "F2F4F7")
    set_cell_margins(c, top=120, bottom=120, left=180, right=180)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(31, 78, 121)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(10.5)
    
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2E75B6"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_usecase_table(doc, table_num, title, uc_id, uc_name, actors, desc, preconditions, requirements, main_flow, sub_flow):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(3)
    r_title = p_title.add_run(f"Bảng {table_num}. Đặc tả Use Case \"{title}\"")
    r_title.bold = True
    r_title.italic = True
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(31, 78, 121)

    table = doc.add_table(rows=8, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.8), Inches(4.7)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    fields = [
        ("Use Case ID", uc_id),
        ("Tên Use Case", uc_name),
        ("Tác nhân", actors),
        ("Mô tả chức năng", desc),
        ("Điều kiện tiên quyết", preconditions),
        ("Yêu cầu", requirements),
        ("Kịch bản chính", main_flow),
        ("Kịch bản phụ", sub_flow)
    ]

    label_bg = "D9E1F2"

    for i, (label, val) in enumerate(fields):
        c_label = table.cell(i, 0)
        c_val = table.cell(i, 1)

        set_cell_background(c_label, label_bg)
        set_cell_margins(c_label, top=70, bottom=70, left=120, right=120)
        set_cell_margins(c_val, top=70, bottom=70, left=120, right=120)

        p_lbl = c_label.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(2)
        p_lbl.paragraph_format.space_after = Pt(2)
        r_lbl = p_lbl.add_run(label)
        r_lbl.bold = True
        r_lbl.font.name = 'Times New Roman'
        r_lbl.font.size = Pt(10.5)

        p_val = c_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(2)
        p_val.paragraph_format.space_after = Pt(2)
        
        lines = val.strip().split("\n")
        for line_idx, line in enumerate(lines):
            if line_idx > 0:
                p_val = c_val.add_paragraph()
                p_val.paragraph_format.space_before = Pt(1)
                p_val.paragraph_format.space_after = Pt(1)
            r_val = p_val.add_run(line.strip())
            r_val.font.name = 'Times New Roman'
            r_val.font.size = Pt(10.5)

    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="8EA9DB"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="8EA9DB"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="8EA9DB"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="8EA9DB"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_c = doc.add_paragraph()
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.paragraph_format.space_before = Pt(4)
        p_c.paragraph_format.space_after = Pt(2)
        r_c = p_c.add_run(caption)
        r_c.bold = True
        r_c.italic = True
        r_c.font.name = 'Times New Roman'
        r_c.font.size = Pt(10)
        r_c.font.color.rgb = RGBColor(80, 80, 80)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    c = table.cell(0, 0)
    c.width = Inches(6.5)
    set_cell_background(c, "F8F9FA")
    set_cell_margins(c, top=80, bottom=80, left=120, right=120)

    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    lines = code_str.strip().split("\n")
    for idx, line in enumerate(lines):
        if idx > 0:
            p = c.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(40, 40, 40)

    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="5B9BD5"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("Generator functions ready.")
